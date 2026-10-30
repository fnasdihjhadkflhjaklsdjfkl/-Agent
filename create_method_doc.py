from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

out = r'D:\ChatGPT\中期\HSMoE_方法部分_中文.docx'
doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.85); sec.bottom_margin = Inches(0.85)
sec.left_margin = Inches(0.9); sec.right_margin = Inches(0.9)
styles = doc.styles
styles['Normal'].font.name = '宋体'; styles['Normal']._element.rPr.rFonts.set(qn('w:eastAsia'),'宋体'); styles['Normal'].font.size = Pt(10.5)
for name, size in [('Title',18),('Heading 1',15),('Heading 2',12.5),('Heading 3',11.5)]:
    s=styles[name]; s.font.name='黑体'; s._element.rPr.rFonts.set(qn('w:eastAsia'),'黑体'); s.font.size=Pt(size); s.font.bold=True

doc.add_paragraph('HSMoE 方法部分', 'Title').alignment = WD_ALIGN_PARAGRAPH.CENTER
p=doc.add_paragraph('基于层次化主导模态选择与 Shapley 路由混合专家的缺失模态多模态情感分析方法')
p.alignment=WD_ALIGN_PARAGRAPH.CENTER
doc.add_paragraph('本章介绍 HSMoE 的整体架构、输入构造、层次化主导模态选择、Shapley 路由混合专家、专家重建以及联合优化目标。')

def h(t,l=1): doc.add_heading(t, level=l)
def para(t): doc.add_paragraph(t)
def eq(t):
    p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.CENTER; r=p.add_run(t); r.font.name='Cambria Math'; r.font.size=Pt(10.5)

h('1 整体架构',1)
para('给定文本、视觉和音频三种模态输入，HSMoE 首先利用模态专属编码器提取高层语义特征；随后，层次化主导模态选择模块（Hierarchical Dominant-modality Selection，HDS）根据局部特征、跨模态交互和模态一致性动态确定当前样本的主导模态；在此基础上，Shapley 路由混合专家模块（Shapley Routing Mixture-of-Experts，SRMoE）对七个异构专家进行贡献度建模和自适应融合；最后，将融合表示输入多层感知机，输出情感预测结果。')
para('模型同时引入专家重建模块（Expert Reconstruction，ER），用于从不完整输入中恢复缺失模态的潜在信息。整体框架由 HDS、SRMoE 和 ER 三个核心部分组成。')

h('2 多模态输入与缺失模态构造',1)
para('设模态集合为 m∈{t,v,a}，其中 t、v 和 a 分别表示文本、视觉和音频。原始模态特征记为：')
eq('Xᵐ ∈ R^(Tᵐ×dᵐ)')
para('本文采用 BERT 编码文本，使用 OpenFace 提取视觉特征，使用 Librosa 提取音频特征。经过投影和位置编码后，各模态被映射到统一维度，并表示为 Hᵐ。为模拟实际场景中的缺失模态，随机缺失视觉和音频特征时采用零向量填充；缺失文本词语使用 [UNK] 标记替换。最终得到含噪或不完整的输入序列 Hᵐ。')

h('3 层次化主导模态选择',1)
h('3.1 模态感知交互',2)
para('HDS 包含模态感知交互和置信度感知选择两个阶段。首先，对每个模态分别使用自注意力模块增强模态内部依赖关系，并通过池化获得模态级表示：')
eq('Hʰᵐ = SelfAttn(Hᵐ)， pᵐ = Pool(Hʰᵐ)')
para('为刻画模态之间的上下文关系，进一步计算任意有序模态对之间的双向跨模态注意力，并将交互结果组织为 3×3 的模态交互矩阵 M：')
eq('Mᵢⱼ = CrossAttn(Hʰⁱ,Hʰʲ)')
para('矩阵对角元素表示模态内部响应，非对角元素表示不同模态之间的交互强度。')
h('3.2 置信度感知选择',2)
para('为了避免仅依据单一指标判断模态可靠性，本文从三个角度计算模态置信度：局部置信度、全局置信度和一致性置信度。局部置信度根据单个模态的池化表示评估；全局置信度根据模态交互矩阵评估；一致性置信度根据各模态的联合表示评估。最终主导性得分为：')
eq('sₘ = λ₁ Cˡ(pₘ) + λ₂ Cᵍ(M) + λ₃ Cᶜᵒⁿˢ(pₐₗₗ)')
para('其中 λ₁、λ₂ 和 λ₃ 为可学习的融合系数。模型选择得分最高的模态作为当前样本的主导模态 m*，并将其作为后续专家划分和路由的依据。')

h('4 Shapley 路由混合专家',1)
h('4.1 异构专家构建',2)
para('模型包含七个异构专家：三个单模态专家 Eₜ、Eᵥ、Eₐ，三个双模态专家 Eₜᵥ、Eₜₐ、Eᵥₐ，以及一个三模态专家 Eₜᵥₐ。每个专家采用由多头自注意力和前馈网络组成的 Transformer 编码器，并通过残差连接和层归一化稳定训练。')
eq('Hʲ′ᵢ = MSA(LN(Hʲ⁻¹ᵢ)) + Hʲ⁻¹ᵢ')
eq('Hʲᵢ = FFN(LN(Hʲ′ᵢ)) + Hʲ′ᵢ')
para('对于多模态专家，将其对应模态的特征序列进行拼接，再输入 Transformer 编码器，最后通过平均池化获得专家表示 Hᵢ。')
h('4.2 专家联盟划分',2)
para('根据 HDS 输出的主导模态 m*，将七个专家划分为主导专家联盟 Cᵈ 和辅助专家联盟 Cᵃ。以视觉模态为主导时，包含视觉模态的专家 Eᵥ、Eᵥₜ、Eᵥₐ 和 Eᵥₜₐ 被划分到主导联盟，其余专家划分到辅助联盟。由于主导联盟最多包含四个专家，因此其 Shapley 值可以精确计算。')
h('4.3 Shapley 值计算与路由',2)
para('对于主导联盟中的专家，采用合作博弈中的 Shapley 值衡量其边际贡献：')
eq('φᵈᵢ = Σ(S⊆Cᵈ) |S|!(|Cᵈ|-|S|-1)! / |Cᵈ|! · [v(S∪{i})-v(S)]')
para('其中 v(S) 表示专家子集 S 的效用函数。对于辅助联盟，为降低计算开销，采用随机采样近似估计其 Shapley 值。随后，将 Shapley 值归一化为专家路由权重：')
eq('wᵈᵢ = φᵈᵢ / Σᵢ∈Cᵈ φᵈᵢ，  wᵃⱼ = φᵃⱼ / Σⱼ∈Cᵃ φᵃⱼ')
para('最终融合表示由主导专家和辅助专家共同构成：')
eq('Ĥ = Σᵢ∈Cᵈ wᵈᵢ Hᵈᵢ + Σⱼ∈Cᵃ wᵃⱼ Hᵃⱼ')
para('融合表示输入由两层全连接层和 ReLU 激活组成的多层感知机，得到最终情感预测值 ŷ。')

h('5 专家重建模块',1)
para('为减轻模态缺失造成的信息损失，本文设计专家重建器 R。重建器由两个 Transformer 层组成，以不完整输入 Hᵐ 为条件，恢复原始模态特征 Xᵐ。重建损失采用均方误差形式：')
eq('Lᴱᴿ = 1/N Σₘ∈{t,v,a} Σⁿ₌₁ᴺ ||xᵐₙ - Rₘ(hᵐₙ)||²')
para('通过最小化重建误差，模型能够学习不同模态之间的相关性，并在某一模态缺失时恢复其潜在语义信息。')

h('6 联合优化目标',1)
para('为使专家路由结果与 HDS 的主导模态判断保持一致，定义辅助约束损失：')
eq('Lₐᵤₓ = Σₘ∈{t,v,a} (sₘ - gₘ)²')
para('其中 sₘ 为 HDS 输出的模态主导性得分，gₘ 为所有包含模态 m 的专家权重之和。情感预测损失采用 L1 损失：')
eq('Lₛₚ = 1/N Σⁿ₌₁ᴺ |ŷₙ-yₙ|')
para('最终总体损失由路由一致性损失、专家重建损失和情感预测损失加权组成：')
eq('L = αLₐᵤₓ + βLᴱᴿ + γLₛₚ')
para('其中 α、β 和 γ 为控制不同训练目标相对贡献的超参数。通过联合优化上述目标，模型能够同时学习主导模态选择、专家贡献分配、缺失信息恢复和情感预测。')

doc.save(out)
print(out)
